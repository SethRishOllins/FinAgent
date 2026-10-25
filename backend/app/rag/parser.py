from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path


class DocumentParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedPage:
    text: str
    page_number: int | None
    section: str | None = None


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    content_type: str
    pages: list[ParsedPage]
    metadata: dict


def _sections(text: str, page_number: int | None) -> list[ParsedPage]:
    current_section: str | None = None
    output: list[ParsedPage] = []
    buffer: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        is_heading = stripped and (stripped.startswith("#") or (len(stripped) < 100 and stripped.isupper()))
        if is_heading:
            if buffer:
                output.append(ParsedPage("\n".join(buffer).strip(), page_number, current_section))
                buffer = []
            current_section = stripped.lstrip("#").strip()
        else:
            buffer.append(line)
    if buffer:
        output.append(ParsedPage("\n".join(buffer).strip(), page_number, current_section))
    return [page for page in output if page.text]


def parse_document(name: str, content: bytes, content_type: str | None = None) -> ParsedDocument:
    suffix = Path(name).suffix.lower()
    normalized_type = content_type or "application/octet-stream"
    if not content:
        raise DocumentParseError("The uploaded document is empty.")
    if suffix in {".txt", ".md"}:
        try:
            text = content.decode("utf-8-sig", errors="strict")
        except UnicodeDecodeError as error:
            raise DocumentParseError("Text documents must contain valid UTF-8 text.") from error
        pages = _sections(text, None)
    elif suffix == ".csv":
        decoded = content.decode("utf-8-sig", errors="strict")
        rows = list(csv.reader(io.StringIO(decoded)))
        if not rows or not rows[0] or not any(cell.strip() for row in rows for cell in row):
            raise DocumentParseError("CSV files must contain at least one row.")
        headers = [cell.strip() or f"column_{index + 1}" for index, cell in enumerate(rows[0])]
        records = [f"Columns: {', '.join(headers)}"]
        for row_number, row in enumerate(rows[1:], start=2):
            values = [row[index].strip() if index < len(row) else "" for index in range(len(headers))]
            records.append(f"Row {row_number}: " + "; ".join(f"{header}: {value}" for header, value in zip(headers, values)))
        pages = [ParsedPage("\n".join(records), None, "CSV data")]
    elif suffix == ".docx":
        if not zipfile.is_zipfile(io.BytesIO(content)):
            raise DocumentParseError("DOCX file is malformed or is not a valid Word document.")
        try:
            from docx import Document as WordDocument
            word = WordDocument(io.BytesIO(content))
            sections: list[ParsedPage] = []
            current_heading: str | None = None
            text: list[str] = []
            for paragraph in word.paragraphs:
                value = paragraph.text.strip()
                if not value:
                    continue
                if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                    if text:
                        sections.append(ParsedPage("\n".join(text), None, current_heading))
                        text = []
                    current_heading = value
                else:
                    text.append(value)
            for table in word.tables:
                for row in table.rows:
                    text.append("; ".join(cell.text.strip() for cell in row.cells))
            if text:
                sections.append(ParsedPage("\n".join(text), None, current_heading))
            pages = sections
            normalized_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except Exception as error:
            raise DocumentParseError("DOCX file could not be parsed safely.") from error
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise DocumentParseError("Encrypted PDFs are not supported.")
            pages = []
            for number, page in enumerate(reader.pages, start=1):
                extracted = (page.extract_text() or "").strip()
                if extracted:
                    pages.extend(_sections(extracted, number))
            normalized_type = "application/pdf"
        except DocumentParseError:
            raise
        except Exception as error:
            raise DocumentParseError("PDF could not be parsed safely; the file may be malformed.") from error
        if not pages:
            raise DocumentParseError("PDF has no extractable text; it may be scanned or image-only and requires OCR. OCR is not configured.")
    else:
        raise DocumentParseError("Supported file formats are PDF, TXT, Markdown, DOCX, and CSV.")
    if not pages or not any(page.text.strip() for page in pages):
        raise DocumentParseError("The document contains no extractable text.")
    return ParsedDocument(title=Path(name).stem, content_type=normalized_type, pages=pages, metadata={"original_filename": name, "page_count": len({page.page_number for page in pages if page.page_number is not None}), "ocr_required": False})
