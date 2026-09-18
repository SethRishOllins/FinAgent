from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "reports" / "FinAgent_Major_Project_Report.docx"
FIG_DIR = ROOT / "docs" / "reports" / "figures"
FIG_DIR.mkdir(parents=True, exist_ok=True)


def font(size, bold=False):
    try:
        return ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def diagram(name, title, columns):
    image = Image.new("RGB", (1600, 720), "white")
    draw = ImageDraw.Draw(image)
    draw.text((800, 35), title, fill="#111111", font=font(34, True), anchor="ma")
    count = len(columns)
    width, height = 260, 175
    gap = (1600 - count * width) // (count + 1)
    positions = []
    for index, (heading, body) in enumerate(columns):
        x = gap + index * (width + gap); y = 250
        positions.append((x, y))
        draw.rounded_rectangle((x, y, x + width, y + height), radius=16, fill="#EAF2F0", outline="#23564B", width=3)
        draw.text((x + width / 2, y + 36), heading, fill="#173B33", font=font(24, True), anchor="ma")
        lines = body.split("\n")
        for line_index, line in enumerate(lines):
            draw.text((x + width / 2, y + 83 + line_index * 29), line, fill="#263630", font=font(18), anchor="ma")
        if index < count - 1:
            sx = x + width + 10; ex = x + width + gap - 10; cy = y + height / 2
            draw.line((sx, cy, ex, cy), fill="#5A786E", width=5)
            draw.polygon([(ex, cy), (ex - 16, cy - 10), (ex - 16, cy + 10)], fill="#5A786E")
    image.save(FIG_DIR / name)


diagram("architecture.png", "FinAgent Proposed System Architecture", [
    ("React Vite SPA", "Home and research pages\nPortfolio and experiments\nChat interface"),
    ("FastAPI API", "Typed request schemas\nCORS enabled routes\nError responses"),
    ("Research Services", "Orchestrator and critic\nAnalytics and workspace\nDemo data providers"),
    ("Outputs", "Research reports\nDocument excerpts\nRisk and experiment metrics"),
])
diagram("workflow.png", "Autonomous Research Workflow", [
    ("Request", "Ticker and research\nquestion"),
    ("Specialist Stages", "Market, financial, news\nRAG, sentiment, risk\ncompetitor analysis"),
    ("Critic", "Marks demo inputs\nFlags missing evidence\nSuppresses claims"),
    ("Synthesis", "Structured findings\nConfidence and limits\nFinancial disclaimer"),
])
diagram("document_flow.png", "Current Document Research Data Flow", [
    ("Upload", "TXT, MD, or CSV\nText decoding"),
    ("Workspace Store", "In-memory document\nname, company, text\ntimestamp"),
    ("Keyword Search", "Token overlap score\nTop three excerpts"),
    ("Grounded Chat", "Evidence excerpts\nNo fabricated sources\nSession-only state"),
])
diagram("quant_flow.png", "Quantitative Analytics and Experiment Flow", [
    ("Synthetic Prices", "Deterministic 90 session\nprice series by ticker"),
    ("Analytics", "RSI, volatility\ndrawdown, VaR\nconcentration"),
    ("Experiments", "Naive and trend forecast\nMA crossover backtest\ntransaction cost"),
    ("Disclosure", "Historical/demo output\nNo prediction claim\nNo investment advice"),
])


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr(); mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None: mar = OxmlElement("w:tcMar"); tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None: node = OxmlElement(f"w:{side}"); mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr(); header = OxmlElement("w:tblHeader"); header.set(qn("w:val"), "true"); tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr(); node = OxmlElement("w:cantSplit"); tr_pr.append(node)


def add_page_field(paragraph):
    run = paragraph.add_run(); fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); run._r.addnext(fld)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.7); section.bottom_margin = Inches(.7); section.left_margin = Inches(.75); section.right_margin = Inches(.75)
styles = doc.styles
styles["Normal"].font.name = "Times New Roman"; styles["Normal"].font.size = Pt(11)
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman"); styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11)):
    style = styles[name]; style.font.name = "Times New Roman"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman"); style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.paragraph_format.space_before = Pt(12); style.paragraph_format.space_after = Pt(6)
footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; footer.add_run("FinAgent Major Project Report | Page "); add_page_field(footer)


def p(text="", bold_lead=None):
    paragraph = doc.add_paragraph(); paragraph.paragraph_format.line_spacing = 1.25; paragraph.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True; paragraph.add_run(text[len(bold_lead):])
    else: paragraph.add_run(text)
    return paragraph


def heading(text, level=1):
    return doc.add_heading(text, level=level)


def caption(text, kind="Figure"):
    paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text); run.bold = True; run.font.size = Pt(10)


def table(headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers)); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.style = "Table Grid"
    hdr = tbl.rows[0]; set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]; cell.text = value; set_cell_shading(cell, "1F4E4D"); set_cell_margins(cell)
        for run in cell.paragraphs[0].runs: run.font.color.rgb = RGBColor(255,255,255); run.font.bold = True; run.font.size = Pt(9)
    for row_index, row in enumerate(rows):
        cells = tbl.add_row().cells
        prevent_row_split(tbl.rows[-1])
        for idx, value in enumerate(row):
            cells[idx].text = str(value); set_cell_margins(cells[idx]); cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1: set_cell_shading(cells[idx], "F2F6F5")
            for run in cells[idx].paragraphs[0].runs: run.font.size = Pt(9)
    if widths:
        for row in tbl.rows:
            for idx, width in enumerate(widths): row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return tbl


def figure(filename, number, title):
    paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(FIG_DIR / filename), width=Inches(6.6))
    caption(f"Figure {number}: {title}")


def page_break(): doc.add_page_break()


# Front matter
for _ in range(5): doc.add_paragraph()
title = doc.add_paragraph(); title.style = "Title"; title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.add_run("FinAgent")
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER; sub.add_run("An Autonomous Financial Research and Decision Support Platform").bold = True; sub.runs[0].font.size = Pt(16)
for _ in range(5): doc.add_paragraph()
for text in ["MAJOR PROJECT REPORT", "Submitted in partial fulfillment of the requirements for the degree of", "Bachelor of Technology", "Computer Science and Engineering", "Submitted by", "Rishabh Kumar Chettiar", "Enrollment Number: ____________________", "Under the guidance of", "Supervisor Name: ____________________", "Department of Computer Science and Engineering and Information Technology", "Jaypee Institute of Information Technology, Noida", "Academic Session: ____________________"]:
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=para.add_run(text); run.font.size=Pt(12 if text not in ("MAJOR PROJECT REPORT","Rishabh Kumar Chettiar") else 14); run.bold=text in ("MAJOR PROJECT REPORT","Rishabh Kumar Chettiar")
page_break()

heading("Certificate", 1)
p("This is to certify that the work titled 'FinAgent: An Autonomous Financial Research and Decision Support Platform' is a bona fide major project report prepared by the student named on the title page under the supervision of the undersigned. The project is submitted in partial fulfillment of the academic requirements of the programme. The institutional fields below are intentionally retained for completion by the department.")
p("Name of Supervisor: ______________________________")
p("Designation: _____________________________________")
p("Signature of Supervisor: __________________________")
p("Date: ____________________________________________")
page_break()

heading("Self Declaration of Originality", 1)
p("I declare that this report describes the FinAgent software project developed for academic purposes. The project-specific application logic, API orchestration, analytics functions, user-interface composition, and documentation were developed for this repository. Open-source frameworks and libraries are acknowledged in the declaration below and in the reference list. No claim is made that externally supplied frameworks, package functionality, or future integrations were written from scratch.")
heading("Declaration Regarding Open Source and Prebuilt Components", 2)
table(["Category", "Components evidenced in the repository", "Use in FinAgent"], [
    ["Frontend", "React, React DOM, Vite, TypeScript, Recharts, Lucide React", "Single-page interface, build tooling, charts/icons."],
    ["Backend", "FastAPI, Uvicorn, Pydantic, python-multipart", "HTTP API, typed validation, multipart uploads."],
    ["Analytics and testing", "NumPy, pytest, httpx", "Numerical analytics and automated tests."],
    ["Planned optional services", "PostgreSQL, Redis, Qdrant, OpenAI, Ollama", "Declared in configuration or documentation; not implemented as running project dependencies."],
], [1.15, 2.8, 2.4])
p("The repository does not provide an objective method for calculating a percentage of original code. Therefore, no numerical code-contribution estimate is asserted in this report.")
page_break()

heading("Acknowledgement", 1)
p("I express gratitude to the faculty members, project supervisor, and peers who provide academic guidance and feedback for this major-project work. I also acknowledge the open-source communities behind the development frameworks and research publications that informed the conceptual background of this prototype.")
page_break()

heading("Abstract", 1)
p("FinAgent is a web-based financial research and decision-support prototype for Indian-equity-oriented analysis. The project addresses the difficulty of combining market context, financial indicators, document evidence, portfolio concentration, and uncertainty disclosures in a single research workspace. It is designed as a research aid and explicitly does not provide investment advice, buy/sell instructions, or guaranteed predictions.")
p("The implemented system uses a React and Vite frontend with a FastAPI backend. Its backend exposes typed APIs for company lookup, research execution, portfolio concentration, document upload and search, report history, market overview, peer comparison, forecast experiments, moving-average backtesting, and a grounded chat interaction. The research workflow records a master orchestration stage followed by market, financial, news, document, sentiment, risk, competitor, critic, and synthesis stages. The critic labels missing evidence and prevents the output from being presented as definitive financial guidance.")
p("The present repository intentionally operates with deterministic offline demonstration data. It includes technical indicators, risk statistics, two simple historical forecast baselines, and a moving-average crossover backtest with transaction cost. It also provides in-memory document storage with keyword-overlap retrieval. Five automated tests passed during the documented verification run. Persistent storage, live market/news providers, PDF extraction, vector embeddings, trained sentiment models, and production LLM providers are not implemented and are documented as future scope.")
p("Keywords: financial research, multi-agent orchestration, FastAPI, React, retrieval-assisted document search, portfolio concentration, backtesting, responsible AI.")
page_break()

heading("List of Figures", 1)
for text in ["Figure 3.1 Proposed System Architecture", "Figure 3.2 Autonomous Research Workflow", "Figure 3.3 Current Document Research Data Flow", "Figure 3.4 Quantitative Analytics and Experiment Flow"]: p(text)
heading("List of Tables", 1)
for text in ["Table 1.1 Project Development Timeline", "Table 2.1 Comparative Literature Survey", "Table 3.1 Functional Requirements", "Table 3.2 Technology Stack", "Table 3.3 Implemented API Endpoints", "Table 4.1 Functional Test Evidence", "Table 5.1 Evidence Based Results", "Table 6.1 Limitations and Future Improvements"]: p(text)
heading("Table of Contents", 1)
for text in ["Chapter 1 Introduction", "Chapter 2 Background Study and Literature Survey", "Chapter 3 System Design and Methodology", "Chapter 4 Implementation and Testing", "Chapter 5 Results and Discussion", "Chapter 6 Conclusion and Future Scope", "References"]: p(text)
page_break()

# Chapter 1
heading("Chapter 1 Introduction", 1)
heading("1.1 Background", 2)
p("Financial research commonly requires information from different forms of evidence: price movement, financial statements, market context, company news, portfolio exposure, and analyst interpretation. A user who attempts to combine these sources manually must also judge freshness, provenance, and uncertainty. Language-model-based interfaces can make such analysis easier to ask for, but language generation alone does not establish that a conclusion is supported by data. FinAgent is designed around this distinction. It combines deterministic backend calculations with a structured orchestration record and visible limitations.")
heading("1.2 Problem Statement", 2)
p("The problem addressed by FinAgent is the absence of a compact research workspace that can present a company or portfolio question together with calculations, evidence boundaries, and a clear indication of what the system could not verify. The project does not attempt to automate investment decisions. Instead, it provides an academic prototype in which analytical modules and an agent-style workflow can be inspected by a user or evaluator.")
heading("1.3 Objectives", 2)
table(["Objective", "Implemented evidence"], [
    ["Provide company research", "Company lookup, overview, financial dataset, technical summary, news sample, peer comparison, and research route."],
    ["Make orchestration visible", "Research response contains agent events, findings, warnings, confidence values, and limitations."],
    ["Support portfolio risk discussion", "Portfolio endpoint calculates largest weight and Herfindahl-style concentration index."],
    ["Support document-grounded interaction", "Text document upload, session workspace, keyword-overlap search, and chat citations."],
    ["Support transparent quantitative experiments", "Forecast baseline and trend model; moving-average backtest with transaction cost and stated limitations."],
], [2.25, 4.1])
heading("1.4 Societal and Environmental Relevance", 2)
p("The platform can improve financial literacy by translating concentration and research-coverage concepts into plain-language UI messages. Its disclosures are socially relevant because financial systems can cause harm when uncertain output is framed as certainty. The prototype encourages users to distinguish historical/model-derived signals from advice. Environmentally, a digital research workspace may reduce paper handling for notes and downloaded reports; however, the project has not measured environmental impact. Future deployment should consider the energy cost of external model inference and data services.")
heading("1.5 Project Development Timeline", 2)
table(["Phase", "Major activities", "Evidence in repository"], [
    ["Planning", "Architecture, API and agent documentation", "docs/ARCHITECTURE.md, docs/API.md, docs/AGENTS.md"],
    ["Core implementation", "React interface and FastAPI routes", "frontend/src and backend/app"],
    ["Analytics", "Technical, risk, forecast, and backtest functions", "backend/app/services/analytics.py"],
    ["Verification", "Unit tests, API smoke checks, frontend build", "backend/tests and PROJECT_STATUS.md"],
    ["Future integration", "Persistent data, live providers, vector retrieval, trained ML", "PROJECT_STATUS.md and .env.example"],
], [1.1, 2.4, 2.85])
heading("1.6 Scope", 2)
p("The scope is a locally runnable demonstration system for a small supported ticker universe: RELIANCE, TCS, INFY, and HDFCBANK. It includes deterministic price data and offline financial/news samples to enable reproducible demonstrations. The scope excludes live-data accuracy, investor suitability, user accounts, permanent data storage, document PDFs, and production cloud deployment.")
page_break()

# Chapter 2
heading("Chapter 2 Background Study and Literature Survey", 1)
heading("2.1 Overview of Existing Work", 2)
p("Retrieval-augmented generation was introduced to combine parametric language models with an external non-parametric memory, improving provenance and factuality for knowledge-intensive tasks [1]. Financial language-model research such as FinGPT highlights the need for curated and accessible financial data rather than relying solely on generic language-model knowledge [2]. Multi-agent surveys describe role differentiation, coordination, and task decomposition as central design concerns in complex LLM-based systems [3]. These ideas motivate FinAgent's separation between calculation services, research stages, critic review, and user-facing synthesis.")
heading("2.2 Comparative Literature Survey", 2)
table(["Work", "Method", "Relevant contribution", "Gap addressed by FinAgent"], [
    ["Lewis et al. (2020) [1]", "Retrieval-augmented generation", "External retrieval can improve factual responses and provenance.", "Prototype uses session document search, but does not yet implement dense retrieval."],
    ["Yang et al. (2023) [2]", "Open financial LLM framework", "Emphasizes financial-data curation and accessible FinLLM resources.", "FinAgent keeps a provider boundary and labels demo data rather than claiming live intelligence."],
    ["Guo et al. (2024) [3]", "LLM multi-agent survey", "Reviews roles, communication, and challenges in multi-agent systems.", "FinAgent makes stages and critic output visible in a research report."],
    ["Sharpe (1994) [4]", "Risk-adjusted return measure", "Defines a widely used return-to-risk interpretation.", "FinAgent computes a Sharpe-like statistic for backtest output, with synthetic-data limitation."],
    ["Wang et al. (2024) [5]", "Autonomous-agent survey", "Systematizes planning, memory, and tool use in LLM agents.", "FinAgent implements bounded orchestration records but not a production LLM agent loop."],
], [1.2, 1.3, 2.25, 2.0])
heading("2.3 Limitations of Existing Approaches", 2)
p("Existing work demonstrates useful components but does not remove the engineering challenge of presenting trustworthy research to an end user. A generic chat interface can hide sources, calculation logic, or missing data. A technical dashboard can expose metrics without helping a non-expert understand what needs attention. FinAgent addresses this presentation gap by combining a plain-language portfolio-health view with an inspectable research workflow. This is a design objective rather than a claim of superior analytical performance.")
heading("2.4 Technical Gap and Motivation", 2)
p("The technical gap selected for this project is the integration of research workflow traceability, calculations, document excerpts, and explicit uncertainty in a lightweight local prototype. FinAgent is motivated by the need to demonstrate an agentic architecture without representing planned services as operational. Its critic stage and data-status labels are therefore core design choices.")
page_break()

# Chapter 3
heading("Chapter 3 System Design and Methodology", 1)
heading("3.1 Proposed System", 2)
p("FinAgent has a single-page React user interface and a versioned FastAPI backend. The frontend invokes API endpoints at localhost during development. The backend validates requests with Pydantic, obtains deterministic market/financial samples, runs calculations through NumPy, and returns JSON structures that preserve findings, confidence values, limitations, and warnings. The project is organized so that future provider and persistence adapters can replace demonstration services without changing the principal API surface.")
figure("architecture.png", "3.1", "Proposed System Architecture")
heading("3.2 Autonomous Research Methodology", 2)
p("A research request contains a supported ticker and a question. The ResearchOrchestrator creates a run identifier and writes a sequence of activity events. It computes technical context, reads the offline fundamentals record, classifies available offline news samples, checks the session document workspace, records sentiment scope, and computes risk context. It then records a competitor-analysis scope, critic findings, and a judge/synthesis activity. The final response includes a disclaimer. The current implementation is a deterministic structured workflow; it does not invoke an LLM provider despite configuration placeholders for OpenAI and Ollama.")
figure("workflow.png", "3.2", "Autonomous Research Workflow")
heading("3.3 Data Flow and Document Research", 2)
p("The document endpoint accepts a multipart text file, decodes it as UTF-8 with replacement behavior, rejects content shorter than 20 characters, and stores the name, optional company label, text content, timestamp, and UUID in an in-memory WorkspaceStore. Search tokenizes the query into words longer than three characters and scores documents by token presence. The top three excerpts are returned. The chat service uses these excerpts when available and otherwise offers an offline research response or a safe refusal for buy/sell requests. This is a keyword retrieval implementation, not the planned Qdrant embedding pipeline.")
figure("document_flow.png", "3.3", "Current Document Research Data Flow")
heading("3.4 Quantitative Methodology", 2)
p("For each supported ticker, DemoMarketDataProvider deterministically derives a random seed from the symbol and creates a 90-session synthetic price series. technical_summary computes returns, a 20-session average, 14-period RSI, annualized historical volatility, and maximum drawdown. risk_profile derives volatility, a simple beta proxy, Sharpe and Sortino-like values, drawdown, one-day 95% VaR, and CVaR from the same series. These outputs are demonstrations of calculation flow, not calibrated real-market estimates.")
p("forecast_experiment compares a naive last-value baseline with a linear trend extrapolation over a five-day horizon. The returned MAE, RMSE, and directional-accuracy values are calculations on the synthetic series. backtest_ma_crossover uses eight- and twenty-period moving averages, uses the previous available signal to prevent direct look-ahead exposure, subtracts a configurable 0.1% transaction cost on signal turnover, and returns equity-curve and summary metrics. XGBoost, LSTM, portfolio optimization, anomaly detection, and trained sentiment classification are not implemented.")
figure("quant_flow.png", "3.4", "Quantitative Analytics and Experiment Flow")
heading("3.5 Functional Requirements", 2)
table(["ID", "Requirement", "Implementation status"], [
    ["FR-01", "Search supported companies", "Implemented through GET /api/v1/companies."],
    ["FR-02", "Run autonomous research", "Implemented through POST /api/v1/research with event trace."],
    ["FR-03", "Analyze portfolio concentration", "Implemented through POST /api/v1/portfolio/analyze."],
    ["FR-04", "Upload and search text documents", "Implemented; state is in memory for the API process."],
    ["FR-05", "Provide chat assistance", "Implemented in grounded offline mode with safety response."],
    ["FR-06", "Run forecast and backtest experiments", "Implemented for deterministic offline price samples."],
    ["FR-07", "Persist users, documents, and reports", "Not implemented; PostgreSQL/Qdrant are future integrations."],
], [.7, 2.35, 3.7])
heading("3.6 API and Technology Design", 2)
table(["Layer", "Technology", "Justification and role"], [
    ["Frontend", "React, TypeScript, Vite", "Component-driven single-page interface and production build pipeline."],
    ["Backend", "FastAPI, Pydantic, Uvicorn", "Typed request validation and local API service."],
    ["Numerical layer", "NumPy", "Vectorized calculations for price returns, drawdown, and experiment metrics."],
    ["Testing", "pytest, httpx", "Automated Python test execution and API-adjacent dependency."],
    ["Optional architecture", "PostgreSQL, Redis, Qdrant, OpenAI/Ollama", "Configured/documented future services; not active in the tested prototype."],
], [1.3, 1.7, 3.75])
heading("3.7 Implemented API Endpoints", 2)
table(["Method", "Endpoint", "Observed responsibility"], [
    ["GET", "/health", "Reports offline-capable service mode and configured LLM provider value."],
    ["GET", "/companies and /companies/{symbol}/overview", "Searches supported tickers and returns samples, technical, financial, risk, and data status."],
    ["POST", "/research", "Runs orchestrator and stores report in session workspace."],
    ["POST", "/portfolio/analyze", "Returns normalized concentration metrics."],
    ["POST", "/chat", "Returns grounded offline assistance and document citations when available."],
    ["POST/GET", "/documents/upload, /documents/text, /documents/search", "Stores and retrieves session documents."],
    ["GET", "/companies/{symbol}/news, /forecast, /backtest", "Returns offline news, forecast, and backtest outputs."],
    ["GET", "/compare and /market-overview", "Returns peer table and offline market context."],
], [.7, 2.5, 3.55])
heading("3.8 Security and Responsible AI Design", 2)
p("The repository keeps secrets out of source through .env exclusion and supplies .env.example. Input models limit string lengths and validate portfolio content through Pydantic. The document upload route rejects short text but does not implement antivirus scanning, file-type verification beyond UI filtering, authentication, authorization, persistence encryption, or rate limiting. The chat service identifies buy/sell phrasing and returns a non-advisory alternative. CORS is limited in code to localhost:5173. These protections are appropriate as local prototype measures but insufficient for production deployment.")
page_break()

# Chapter 4
heading("Chapter 4 Implementation and Testing", 1)
heading("4.1 Development Setup", 2)
p("The repository targets Windows development. The frontend package scripts use Vite for development and production build. The backend requirements specify FastAPI, Uvicorn, Pydantic, NumPy, pytest, httpx, and python-multipart. The documented verification environment used Python 3.13 after dependency constraints were updated for compatible binary packages. Docker Compose defines optional PostgreSQL, Redis, and Qdrant containers, but these services are not connected to the current application logic.")
heading("4.2 Frontend Implementation", 2)
p("The frontend is implemented as a React TypeScript single-page application in frontend/src/main.tsx. It maintains a local page state for Home, Research, Market, Portfolio, Compare, Forecast, Backtest, Documents, and Reports. The interface sends fetch requests to the FastAPI base URL and displays API results or errors. The home view emphasizes portfolio health and data disclosures; the research view exposes agent activity and critic warnings; the document view accepts TXT, MD, and CSV selections; the chat widget sends messages to the backend.")
heading("4.3 Backend Implementation", 2)
p("The FastAPI application in backend/app/main.py wires each route to service functions. Market samples are produced by DemoMarketDataProvider. Financial values are contained in a static CompanyFundamentals dataclass map. WorkspaceStore uses Python dataclasses with lists for documents and reports. The services module contains numerical algorithms, while ResearchOrchestrator constructs typed Finding and AgentEvent objects. The design demonstrates separation of concerns, although production repository interfaces and migrations are not present.")
heading("4.4 Testing Strategy", 2)
p("The repository contains two pytest modules. test_analytics.py verifies technical-summary bounds, concentration classification, and that the moving-average backtest produces a positive final value and non-empty equity curve for a monotonic input. test_workspace.py verifies keyword document matching and verifies that the assistant refuses a buy/sell request. The recorded verification also used FastAPI TestClient smoke calls for overview, forecast, backtest, comparison, market overview, and research endpoints. Five pytest tests passed. No test coverage percentage, browser end-to-end suite, load test, or database integration test is available.")
heading("4.5 Functional Test Evidence", 2)
table(["Test ID", "Scenario", "Evidence", "Observed result", "Status"], [
    ["TC-01", "Technical summary", "test_technical_summary_has_expected_metrics", "Trend and RSI bounds asserted.", "Passed"],
    ["TC-02", "Portfolio concentration", "test_portfolio_flags_concentration", "80/20 allocation classified high concentration.", "Passed"],
    ["TC-03", "Backtest output", "test_backtest_returns_metrics_and_equity_curve", "Positive final value and non-empty curve asserted.", "Passed"],
    ["TC-04", "Document matching", "test_workspace_returns_matching_document_excerpt", "Matching source document returned.", "Passed"],
    ["TC-05", "Advice safety", "test_assistant_refuses_buy_sell_advice", "Refusal string asserted.", "Passed"],
    ["TC-06", "Frontend compilation", "npm run build recorded in PROJECT_STATUS.md", "Vite production build succeeded.", "Passed"],
    ["TC-07", "PDF upload/RAG vectors", "No test or implementation", "Not implemented.", "Not applicable"],
], [.65, 1.4, 2.25, 1.55, .65])
heading("4.6 Reproducibility Notes", 2)
p("The market provider intentionally uses a ticker-derived seed and NumPy random generator. This makes the 90-session demonstration series reproducible for the same symbol in a compatible environment. In-memory report and document state, however, is reset when the API process restarts. Reproducibility therefore applies to the analytical sample generation rather than to durable user-workspace storage.")
page_break()

# Chapter 5
heading("Chapter 5 Results and Discussion", 1)
heading("5.1 Result Analysis", 2)
p("The verified result is functional integration of the local prototype, not financial predictive accuracy. The project status records a successful frontend production build and five passing backend tests. FastAPI smoke checks returned HTTP 200 for company overview, forecast, backtest, comparison, market overview, and autonomous research. These checks demonstrate that the inspected routes can execute on the documented local environment.")
heading("5.2 Evidence Based Results", 2)
table(["Area", "Observed output", "Interpretation boundary"], [
    ["Company research", "Structured findings and agent events for supported symbols.", "Inputs are offline samples; output is research context, not advice."],
    ["Technical/risk analytics", "RSI, volatility, drawdown, VaR/CVaR, and related values returned.", "Values are derived from deterministic synthetic series."],
    ["Forecast experiment", "Naive and trend statistical models return MAE, RMSE, directional accuracy, and a five-day vector.", "Not a trained forecasting benchmark; no real-data predictive claim."],
    ["Backtest", "Equity curve, CAGR, volatility, Sharpe-like score, drawdown, and win rate returned.", "Historical offline experiment; no market strategy validation claim."],
    ["Document workspace", "Top keyword-overlap excerpts can be returned to chat.", "No PDF parsing, dense retrieval, reranking, or persistent citation pages."],
], [1.35, 2.85, 2.15])
heading("5.3 Discussion", 2)
p("The main contribution of the current prototype is traceability. Rather than outputting a single opaque response, the research endpoint includes specialist-stage labels, individual findings, evidence lists, confidence values, limitations, warnings, and a disclaimer. The critic deliberately reports the absence of live news and uploaded filing evidence. This improves the demonstrability of the architecture and provides a safer basis for future integration work.")
p("The prototype does not support claims about model quality, investor returns, real-time accuracy, or retrieval quality. Its forecast values and backtest metrics are useful for explaining system flow and formulas, but they cannot be generalized because their input series is synthetic. Similarly, the financial records are static offline demonstration values. The report therefore treats outputs as functional examples rather than empirical financial results.")
heading("5.4 Evaluation Readiness", 2)
p("The codebase contains extension points for later research evaluation but no executed RAG relevance experiment, hallucination-rate experiment, multi-agent versus single-agent comparison, latency distribution, cost measurement, or ML model comparison. A valid future evaluation must use a recorded dataset, a fixed test protocol, clearly defined baselines, and retained outputs. These requirements are explicitly distinguished from current results.")
page_break()

# Chapter 6
heading("Chapter 6 Conclusion and Future Scope", 1)
heading("6.1 Summary of Work", 2)
p("FinAgent demonstrates a local financial research and decision-support prototype with React frontend pages, FastAPI APIs, deterministic analytical data, a visible multi-stage research workflow, critic warnings, document text search, safe chat behavior, portfolio concentration analysis, peer comparison, forecast baselines, and a moving-average backtest. The work is intentionally bounded as an academic prototype and exposes its data and model limitations in code and UI text.")
heading("6.2 Key Contributions", 2)
p("The project contributes an inspectable research response schema, a deterministic offline provider suitable for demonstrations, a modular service layout, explicit critic output, safe non-advisory chatbot behavior, and verifiable baseline test evidence. The separation between implemented computation and planned provider/model capabilities is a deliberate design contribution because it prevents the user interface from misrepresenting current functionality.")
heading("6.3 Social Impact, Ethics, and Compliance", 2)
p("Financial interfaces can influence user behavior even when they do not execute trades. FinAgent therefore labels itself as research and decision support rather than an advisor. Its assistant does not answer buy/sell requests with a recommendation. The platform should not be used for personal financial suitability analysis without validated data, human oversight, and applicable regulatory review. Future deployments must implement consent, authentication, audit trails, secure storage, rate limits, provider terms compliance, and clear mechanisms for correcting stale or incorrect information.")
heading("6.4 Environmental Sustainability", 2)
p("The local prototype has modest infrastructure needs and can run without GPU computation or large local language models. Optional external containers and future model services would increase resource use. A future deployment should prefer small retrieval models, caching, batched ingestion, and measured resource monitoring rather than treating AI inference as cost-free.")
heading("6.5 Limitations and Future Improvements", 2)
table(["Current limitation", "Why it matters", "Future improvement"], [
    ["Offline demo market, financial, and news data", "No current-market reliability or source freshness.", "Implement compliant live-provider adapters with metadata and rate limits."],
    ["In-memory workspace", "Documents and reports disappear when the server restarts.", "Add SQLAlchemy/PostgreSQL for records and Qdrant for vector collections."],
    ["Text-only upload", "Annual-report PDFs and page citations are unsupported.", "Add PDF extraction, chunk metadata, embeddings, reranking, and page references."],
    ["No production LLM provider", "Agent stages are deterministic, not LLM tool-calling agents.", "Implement OpenAI/Ollama providers with budgets, tracing, and fallback behavior."],
    ["No trained ML models", "Sentiment, anomaly, and advanced forecasting remain absent.", "Create reproducible datasets, training scripts, validation splits, and metrics."],
    ["No authentication or durable alerts", "Not ready for multi-user or production use.", "Implement users, roles, encrypted secrets, notifications, and audit logging."],
], [1.55, 2.25, 2.55])
heading("6.6 Conclusion", 2)
p("FinAgent establishes a credible foundation for a larger major-project platform by integrating user-facing financial research flows with backend calculation and evidence-bound orchestration. The project’s strongest current property is that it shows what it knows, what it calculated, and what remains unavailable. Advancing the project should prioritize trustworthy data ingestion and durable evidence storage before adding larger models or more ambitious prediction claims.")
page_break()

heading("References", 1)
references = [
    "[1] P. Lewis et al., 'Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks,' Advances in Neural Information Processing Systems, vol. 33, 2020. https://arxiv.org/abs/2005.11401",
    "[2] H. Yang, X.-Y. Liu, and C. D. Wang, 'FinGPT: Open-Source Financial Large Language Models,' arXiv:2306.06031, 2023. https://arxiv.org/abs/2306.06031",
    "[3] T. Guo et al., 'Large Language Model based Multi-Agents: A Survey of Progress and Challenges,' arXiv:2402.01680, 2024. https://arxiv.org/abs/2402.01680",
    "[4] W. F. Sharpe, 'The Sharpe Ratio,' The Journal of Portfolio Management, vol. 21, no. 1, pp. 49-58, 1994. https://web.stanford.edu/~wfsharpe/art/sr/sr.htm",
    "[5] L. Wang et al., 'A Survey on Large Language Model based Autonomous Agents,' arXiv:2308.11432, 2024. https://arxiv.org/abs/2308.11432",
    "[6] FastAPI, 'FastAPI Documentation.' https://fastapi.tiangolo.com/ (accessed September 2026).",
    "[7] React, 'React Documentation.' https://react.dev/ (accessed September 2026).",
    "[8] NumPy Developers, 'NumPy Documentation.' https://numpy.org/doc/ (accessed September 2026).",
    "[9] FinAgent source repository, local project files reviewed: frontend/src, backend/app, backend/tests, docs, and configuration files, September 2026.",
]
for ref in references: p(ref)

doc.core_properties.title = "FinAgent Major Project Report"
doc.core_properties.author = "Rishabh Kumar Chettiar"
doc.core_properties.subject = "Academic report for FinAgent financial research prototype"
doc.save(OUT)
print(OUT)
